from __future__ import annotations
import base64, hashlib, hmac, json, os, time
from collections import defaultdict, deque
from datetime import datetime, timezone
from io import BytesIO
from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from db import connect, init_db, password_hash

SECRET=os.getenv('JWT_SECRET','dmt-demo-secret-change-in-production').encode(); JWT_TTL=60*60*8
LISTENER_API_KEY=os.getenv('LISTENER_API_KEY','change-this-listener-key')

def b64(v:bytes)->str:return base64.urlsafe_b64encode(v).rstrip(b'=').decode()
def token(username,role):
 h=b64(json.dumps({'alg':'HS256','typ':'JWT'},separators=(',',':')).encode()); p=b64(json.dumps({'sub':username,'role':role,'exp':int(time.time())+JWT_TTL},separators=(',',':')).encode()); s=b64(hmac.new(SECRET,f'{h}.{p}'.encode(),hashlib.sha256).digest()); return f'{h}.{p}.{s}'
def decode(t):
 try:
  h,p,s=t.split('.'); expected=b64(hmac.new(SECRET,f'{h}.{p}'.encode(),hashlib.sha256).digest()); data=json.loads(base64.urlsafe_b64decode(p+'='*(-len(p)%4)))
  if not hmac.compare_digest(s,expected) or data['exp']<time.time(): raise ValueError()
  return data
 except Exception as e: raise HTTPException(401,'Token inválido o expirado') from e
def current_user(request:Request):
 a=request.headers.get('Authorization','');
 if not a.startswith('Bearer '): raise HTTPException(401,'Autenticación requerida')
 u=decode(a[7:])
 return u
def current_admin(request:Request):
 u=current_user(request)
 if u.get('role')!='ADMINISTRATOR': raise HTTPException(403,'Se requiere rol administrador')
 return u
def attendance_writer(request:Request):
    listener_key=request.headers.get('X-Listener-Key','')
    if listener_key and hmac.compare_digest(listener_key,LISTENER_API_KEY): return {'role':'LISTENER'}
    return current_admin(request)
class Login(BaseModel): username:str; password:str
class StudentIn(BaseModel): biometric_id:str=Field(min_length=1,max_length=64); first_name:str=Field(min_length=1,max_length=160); last_name:str=Field(min_length=1,max_length=160); course_id:int=Field(ge=1,le=13); email:str|None=None; active:bool=True
class StudentResetIn(BaseModel): confirmation:str
class InspectorIn(BaseModel): first_name:str; last_name:str; username:str=Field(min_length=4); email:str|None=None; password:str='change-me'; type:str='CURSO'; course_ids:list[int]=[]; active:bool=True
class ScheduleIn(BaseModel): course_id:int=Field(ge=1,le=13); expected_entry:str=Field(pattern=r'^([01]\d|2[0-3]):[0-5]\d$'); late_after:str=Field(pattern=r'^([01]\d|2[0-3]):[0-5]\d$')
app=FastAPI(title='DMT Biometría',version='2.0.0'); app.add_middleware(CORSMiddleware,allow_origins=['http://localhost:5173'],allow_credentials=True,allow_methods=['*'],allow_headers=['*'])
rate_buckets=defaultdict(deque)
@app.on_event('startup')
def startup(): init_db()
@app.middleware('http')
async def rate_limit(request,call_next):
 key=f"{request.client.host if request.client else 'unknown'}:{request.url.path}:{request.method}"; now=time.time(); bucket=rate_buckets[key]
 while bucket and bucket[0]<now-60: bucket.popleft()
 if len(bucket)>=(10 if request.url.path=='/api/auth/login' else 120): raise HTTPException(429,'Demasiadas solicitudes. Intente nuevamente en un minuto.')
 bucket.append(now); return await call_next(request)
@app.get('/api/health')
def health(): return {'status':'ok','database':'postgresql'}
@app.post('/api/auth/login')
def login(data:Login):
 with connect() as c: u=c.execute('SELECT username,role,password_hash FROM app_users WHERE username=%s AND active', (data.username,)).fetchone()
 if not u or not hmac.compare_digest(u['password_hash'],password_hash(data.password)): raise HTTPException(401,'Usuario o contraseña incorrectos')
 return {'token':token(u['username'],u['role']),'user':{'username':u['username'],'role':u['role']}}
@app.get('/api/auth/me')
def me(user=Depends(current_user)): return {'username':user['sub'],'role':user['role']}
@app.get('/api/courses')
def courses(user=Depends(current_admin)):
 with connect() as c:return c.execute('SELECT id,name,active FROM courses WHERE active ORDER BY id').fetchall()
def student_row(r): return {**r,'course':r.pop('course_name')}
@app.get('/api/students')
def list_students(search:str='',course_id:int|None=None,page:int=1,page_size:int=25,sort_by:str='last_name',sort_dir:str='asc',user=Depends(current_admin)):
 page=max(1,page); page_size=min(100,max(1,page_size)); order='biometric_id' if sort_by=='biometric_id' else 'last_names'; direction='DESC' if sort_dir=='desc' else 'ASC'; clauses=['s.active']; params=[]
 if course_id: clauses.append('s.course_id=%s'); params.append(course_id)
 if search: clauses.append("(s.first_names||' '||s.last_names||' '||s.biometric_id||' '||COALESCE(s.representative_email,'')) ILIKE %s"); params.append(f'%{search}%')
 where=' AND '.join(clauses)
 with connect() as c:
  total=c.execute(f'SELECT count(*) AS n FROM students s WHERE {where}',params).fetchone()['n']; params_page=params+[page_size,(page-1)*page_size]
  rows=c.execute(f'''SELECT s.id,s.biometric_id,s.first_names AS first_name,s.last_names AS last_name,s.course_id,s.representative_email AS email,c.name AS course,s.active FROM students s JOIN courses c ON c.id=s.course_id WHERE {where} ORDER BY s.{order} {direction},s.id LIMIT %s OFFSET %s''',params_page).fetchall()
 return {'items':[dict(r) for r in rows],'page':page,'page_size':page_size,'total':total,'pages':max(1,(total+page_size-1)//page_size)}
@app.post('/api/students',status_code=201)
def create_student(data:StudentIn,user=Depends(current_admin)):
 try:
  with connect() as c:
   r=c.execute('''INSERT INTO students(biometric_id,first_names,last_names,course_id,representative_email,active) VALUES(%s,%s,%s,%s,%s,%s) RETURNING id''',(data.biometric_id.strip(),data.first_name.strip(),data.last_name.strip(),data.course_id,data.email or None,data.active)).fetchone(); c.commit(); return {'id':r['id'],**data.model_dump()}
 except Exception as e: raise HTTPException(409,'El ID biométrico ya existe o los datos no son válidos') from e
@app.put('/api/students/{student_id}')
def update_student(student_id:int,data:StudentIn,user=Depends(current_admin)):
 with connect() as c:
  r=c.execute('''UPDATE students SET biometric_id=%s,first_names=%s,last_names=%s,course_id=%s,representative_email=%s,active=%s,updated_at=now() WHERE id=%s RETURNING id''',(data.biometric_id.strip(),data.first_name.strip(),data.last_name.strip(),data.course_id,data.email or None,data.active,student_id)).fetchone();
  if not r: raise HTTPException(404,'Estudiante no encontrado')
  c.commit(); return {'id':student_id,**data.model_dump()}
@app.delete('/api/students/{student_id}')
def delete_student(student_id:int,user=Depends(current_admin)):
 with connect() as c: r=c.execute('UPDATE students SET active=false,updated_at=now() WHERE id=%s RETURNING id',(student_id,)).fetchone(); c.commit()
 if not r: raise HTTPException(404,'Estudiante no encontrado')
 return {'deleted':True}
@app.delete('/api/students')
def clear_students(data:StudentResetIn,user=Depends(current_admin)):
 if data.confirmation.strip().upper()!='ELIMINAR ESTUDIANTES': raise HTTPException(400,'Escriba ELIMINAR ESTUDIANTES para confirmar esta operación')
 with connect() as c:
  notifications=c.execute('DELETE FROM attendance_notifications').rowcount
  attendance=c.execute('DELETE FROM daily_attendance').rowcount
  events=c.execute('DELETE FROM attendance_events').rowcount
  students=c.execute('DELETE FROM students').rowcount
  imports=c.execute('DELETE FROM import_runs').rowcount
  c.commit()
 return {'deleted_students':students,'deleted_attendance':attendance,'deleted_events':events,'deleted_notifications':notifications,'deleted_imports':imports}
@app.get('/api/inspectors')
def list_inspectors(user=Depends(current_admin)):
 with connect() as c:
  rows=c.execute('''SELECT i.id,i.first_names AS first_name,i.last_names AS last_name,u.username,i.email,
      COALESCE(MAX(a.scope::text),'CURSO') AS type,
      COALESCE(array_agg(a.course_id ORDER BY a.course_id) FILTER (WHERE a.course_id IS NOT NULL),'{}') AS course_ids
      FROM inspectors i JOIN app_users u ON u.id=i.user_id
      LEFT JOIN inspector_assignments a ON a.inspector_id=i.id AND a.active
      WHERE i.active GROUP BY i.id,u.username ORDER BY i.last_names,i.first_names''').fetchall()
 return [dict(r,course_ids=list(range(1,14)) if r['type']=='GENERAL' else list(r['course_ids'])) for r in rows]
@app.post('/api/students/import')
async def import_students(file:UploadFile=File(...),mode:str='append',user=Depends(current_admin)):
 if not file.filename or not file.filename.lower().endswith(('.xlsx','.xlsm')): raise HTTPException(400,'Suba un archivo Excel .xlsx')
 from openpyxl import load_workbook
 rows=list(load_workbook(BytesIO(await file.read()),read_only=True,data_only=True).active.iter_rows(values_only=True)); header=next(i for i,r in enumerate(rows) if r and str(r[0]).strip().upper()=='ID'); headers=[str(x).strip().lower() if x else '' for x in rows[header]]; idx={h:i for i,h in enumerate(headers)}; required=['id','nombre','apellido','id de departamento','email']; missing=[x for x in required if x not in idx]
 if missing: raise HTTPException(400,f'Faltan columnas: {", ".join(missing)}')
 parsed=[]; errors=[]; seen=set()
 for n,row in enumerate(rows[header+1:],header+2):
  if not any(v is not None for v in row): continue
  bid=str(row[idx['id']]).strip() if row[idx['id']] is not None else ''
  try: cid=int(row[idx['id de departamento']])
  except: cid=0
  if not bid or bid in seen or cid not in range(1,14): errors.append({'row':n,'message':'ID biométrico vacío/repetido o curso inválido'}); continue
  seen.add(bid); parsed.append((bid,str(row[idx['nombre']] or '').strip(),str(row[idx['apellido']] or '').strip(),cid,str(row[idx['email']]).strip() if row[idx['email']] else None))
 if errors and mode=='replace': return {'inserted':0,'errors':errors,'mode':mode}
 with connect() as c:
  if mode=='replace': c.execute('UPDATE students SET active=false,updated_at=now() WHERE active')
  inserted=0
  for item in parsed:
   try: c.execute('''INSERT INTO students(biometric_id,first_names,last_names,course_id,representative_email) VALUES(%s,%s,%s,%s,%s) ON CONFLICT(biometric_id) DO UPDATE SET first_names=EXCLUDED.first_names,last_names=EXCLUDED.last_names,course_id=EXCLUDED.course_id,representative_email=EXCLUDED.representative_email,active=true,updated_at=now()''',item); inserted+=1
   except Exception: errors.append({'message':f'No se pudo insertar {item[0]}'})
  c.execute('INSERT INTO import_runs(filename,mode,inserted_count,error_count,errors) VALUES(%s,%s,%s,%s,%s)',(file.filename,mode,inserted,len(errors),json.dumps(errors))); c.commit()
 return {'inserted':inserted,'errors':errors,'mode':mode}
@app.post('/api/inspectors',status_code=201)
def create_inspector(data:InspectorIn,user=Depends(current_admin)):
    if data.type not in ('CURSO','GENERAL'): raise HTTPException(400,'Tipo de inspector inválido')
    ids=sorted(set(data.course_ids))
    if data.type=='CURSO' and not ids: raise HTTPException(400,'Seleccione al menos un curso')
    if any(course_id not in range(1,14) for course_id in ids): raise HTTPException(400,'Curso inválido')
    with connect() as c:
        try:
            u=c.execute("INSERT INTO app_users(username,password_hash,role) VALUES(%s,%s,'INSPECTOR') RETURNING id",(data.username,password_hash(data.password or 'change-me'))).fetchone()
            i=c.execute("INSERT INTO inspectors(user_id,first_names,last_names,email) VALUES(%s,%s,%s,%s) RETURNING id",(u['id'],data.first_name,data.last_name,data.email or None)).fetchone()
            if data.type=='GENERAL': c.execute("INSERT INTO inspector_assignments(inspector_id,scope,course_id) VALUES(%s,'GENERAL',NULL)",(i['id'],))
            else:
                for course_id in ids: c.execute("INSERT INTO inspector_assignments(inspector_id,scope,course_id) VALUES(%s,'CURSO',%s)",(i['id'],course_id))
            c.commit(); return {'id':i['id'],**data.model_dump(exclude={'password'}),'course_ids':ids}
        except Exception as e: c.rollback(); raise HTTPException(409,'El usuario o la asignación ya existe') from e
@app.put('/api/inspectors/{inspector_id}')
def update_inspector(inspector_id:int,data:InspectorIn,user=Depends(current_admin)):
    ids=sorted(set(data.course_ids))
    if data.type not in ('CURSO','GENERAL') or (data.type=='CURSO' and not ids) or any(course_id not in range(1,14) for course_id in ids): raise HTTPException(400,'Asignación de cursos inválida')
    with connect() as c:
        i=c.execute('SELECT user_id FROM inspectors WHERE id=%s AND active',(inspector_id,)).fetchone()
        if not i: raise HTTPException(404,'Inspector no encontrado')
        c.execute('UPDATE inspectors SET first_names=%s,last_names=%s,email=%s,updated_at=now() WHERE id=%s',(data.first_name,data.last_name,data.email or None,inspector_id))
        if data.password: c.execute('UPDATE app_users SET username=%s,password_hash=%s,updated_at=now() WHERE id=%s',(data.username,password_hash(data.password),i['user_id']))
        else: c.execute('UPDATE app_users SET username=%s,updated_at=now() WHERE id=%s',(data.username,i['user_id']))
        c.execute('UPDATE inspector_assignments SET active=false WHERE inspector_id=%s',(inspector_id,))
        if data.type=='GENERAL': c.execute("INSERT INTO inspector_assignments(inspector_id,scope,course_id) VALUES(%s,'GENERAL',NULL)",(inspector_id,))
        else:
            for course_id in ids: c.execute("INSERT INTO inspector_assignments(inspector_id,scope,course_id) VALUES(%s,'CURSO',%s)",(inspector_id,course_id))
        c.commit()
    return {'id':inspector_id,**data.model_dump(exclude={'password'}),'course_ids':ids}
@app.delete('/api/inspectors/{inspector_id}')
def delete_inspector(inspector_id:int,user=Depends(current_admin)):
    with connect() as c:
        r=c.execute('UPDATE inspectors SET active=false,updated_at=now() WHERE id=%s RETURNING id',(inspector_id,)).fetchone()
        c.execute('UPDATE inspector_assignments SET active=false WHERE inspector_id=%s',(inspector_id,))
        c.execute('UPDATE app_users SET active=false,updated_at=now() WHERE id=(SELECT user_id FROM inspectors WHERE id=%s)',(inspector_id,)); c.commit()
    if not r: raise HTTPException(404,'Inspector no encontrado')
    return {'deleted':True}
@app.get('/api/schedules')
def list_schedules(user=Depends(current_admin)):
    with connect() as c:
        rows=c.execute("""SELECT c.id AS course_id,c.name,
          COALESCE(MIN(s.expected_entry),'07:00:00'::time) AS expected_entry,
          COALESCE(MIN(s.late_after),'07:15:00'::time) AS late_after
          FROM courses c LEFT JOIN course_schedules s ON s.course_id=c.id AND s.active
          WHERE c.active GROUP BY c.id ORDER BY c.id""").fetchall()
    return [dict(r,expected_entry=str(r['expected_entry'])[:5],late_after=str(r['late_after'])[:5]) for r in rows]
@app.put('/api/schedules/{course_id}')
def update_schedule(course_id:int,data:ScheduleIn,user=Depends(current_admin)):
    if data.course_id!=course_id or data.expected_entry>=data.late_after: raise HTTPException(400,'La hora de atraso debe ser posterior a la hora de entrada')
    with connect() as c:
        for weekday in range(1,8):
            c.execute("""INSERT INTO course_schedules(course_id,weekday,expected_entry,late_after)
              VALUES(%s,%s,%s,%s) ON CONFLICT(course_id,weekday) DO UPDATE SET expected_entry=EXCLUDED.expected_entry,late_after=EXCLUDED.late_after,active=true""",(course_id,weekday,data.expected_entry,data.late_after))
        c.commit()
    return {'course_id':course_id,'expected_entry':data.expected_entry,'late_after':data.late_after}
@app.get('/api/inspector/courses')
def inspector_courses(user=Depends(current_user)):
    if user.get('role')!='INSPECTOR': raise HTTPException(403,'Se requiere rol inspector')
    with connect() as c:
        rows=c.execute("""SELECT c.id,c.name FROM app_users u JOIN inspectors i ON i.user_id=u.id
          JOIN inspector_assignments a ON a.inspector_id=i.id AND a.active JOIN courses c ON c.id=a.course_id
          WHERE u.username=%s AND i.active ORDER BY c.id""",(user['sub'],)).fetchall()
        general=c.execute("""SELECT 1 FROM app_users u JOIN inspectors i ON i.user_id=u.id JOIN inspector_assignments a ON a.inspector_id=i.id AND a.active AND a.scope='GENERAL' WHERE u.username=%s AND i.active""",(user['sub'],)).fetchone()
        if general: rows=c.execute('SELECT id,name FROM courses WHERE active ORDER BY id').fetchall()
    return rows
@app.get('/api/inspector/attendance/today')
def inspector_attendance_today(date:str|None=None,course_id:int|None=None,user=Depends(current_user)):
    if user.get('role')!='INSPECTOR': raise HTTPException(403,'Se requiere rol inspector')
    target=date or datetime.now(timezone.utc).date().isoformat(); params=[user['sub'],target]; course_clause=''
    if course_id: course_clause=' AND s.course_id=%s'; params.append(course_id)
    with connect() as c:
        rows=c.execute(f"""SELECT s.id,s.biometric_id,s.first_names AS first_name,s.last_names AS last_name,
          s.course_id,c.name AS course,d.status,d.first_entry_at,d.last_exit_at
          FROM app_users u JOIN inspectors i ON i.user_id=u.id JOIN inspector_assignments a ON a.inspector_id=i.id AND a.active
          JOIN students s ON s.active AND (a.scope='GENERAL' OR a.course_id=s.course_id)
          JOIN courses c ON c.id=s.course_id LEFT JOIN daily_attendance d ON d.student_id=s.id AND d.attendance_date=%s
          WHERE u.username=%s{course_clause} ORDER BY c.id,s.last_names,s.first_names""",[target,user['sub']]+([course_id] if course_id else [])).fetchall()
    return [dict(r,status=r['status'] or 'ABSENT') for r in rows]
@app.get('/api/attendance/summary')
def attendance_summary(course_id:int|None=None,from_date:str|None=None,to_date:str|None=None,user=Depends(current_admin)):
    clauses=['s.active']; where_params=[]; join_params=[]
    if course_id: clauses.append('s.course_id=%s'); where_params.append(course_id)
    date_clause=''
    if from_date: date_clause+=' AND d.attendance_date >= %s'; join_params.append(from_date)
    if to_date: date_clause+=' AND d.attendance_date <= %s'; join_params.append(to_date)
    with connect() as c:
        rows=c.execute(f'''SELECT s.id,s.biometric_id,s.first_names AS first_name,s.last_names AS last_name,s.course_id,d.attendance_date,d.status,d.first_entry_at,d.last_exit_at FROM students s LEFT JOIN daily_attendance d ON d.student_id=s.id{date_clause} WHERE {' AND '.join(clauses)} ORDER BY s.last_names,s.first_names,d.attendance_date''',join_params+where_params).fetchall()
    return [dict(r) for r in rows]
@app.post('/api/attendance/events',status_code=201)
def create_attendance_event(payload:dict,user=Depends(attendance_writer)):
    biometric_id=str(payload.get('biometric_id','')).strip(); occurred_at=payload.get('occurred_at'); event_type=payload.get('event_type','UNKNOWN')
    if not biometric_id or not occurred_at: raise HTTPException(400,'biometric_id y occurred_at son obligatorios')
    if event_type not in ('ENTRY','EXIT','UNKNOWN'): raise HTTPException(400,'event_type inválido')
    try:
        from zoneinfo import ZoneInfo
        local_dt=datetime.fromisoformat(str(occurred_at).replace('Z','+00:00'))
        local_dt=local_dt.replace(tzinfo=ZoneInfo('America/Guayaquil')) if local_dt.tzinfo is None else local_dt.astimezone(ZoneInfo('America/Guayaquil'))
    except ValueError as e: raise HTTPException(400,'occurred_at debe ser una fecha ISO-8601 válida') from e
    with connect() as c:
        s=c.execute('SELECT id,course_id FROM students WHERE biometric_id=%s AND active',(biometric_id,)).fetchone()
        raw=c.execute('''INSERT INTO attendance_events(student_id,biometric_id,event_type,occurred_at,device_id,source_event_id,raw_payload)
          VALUES(%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (device_id,source_event_id) DO NOTHING RETURNING id''',
          (s['id'] if s else None,biometric_id,event_type,occurred_at,payload.get('device_id'),payload.get('source_event_id'),json.dumps(payload))).fetchone()
        canonical=False; notification_queued=False
        if s and event_type in ('ENTRY','EXIT'):
            day=local_dt.date()
            daily=c.execute('SELECT first_entry_at,last_exit_at,status FROM daily_attendance WHERE attendance_date=%s AND student_id=%s FOR UPDATE',(day,s['id'])).fetchone()
            canonical=daily is None or (daily['first_entry_at'] is None if event_type=='ENTRY' else daily['last_exit_at'] is None)
            status='PRESENT'
            if event_type=='ENTRY':
                schedule=c.execute('SELECT late_after FROM course_schedules WHERE course_id=%s AND weekday=%s AND active',(s['course_id'],local_dt.isoweekday())).fetchone()
                if schedule and local_dt.time()>schedule['late_after']: status='LATE'
                c.execute("""INSERT INTO daily_attendance(attendance_date,student_id,first_entry_at,status)
                  VALUES(%s,%s,%s,%s) ON CONFLICT (attendance_date,student_id) DO UPDATE SET
                    first_entry_at=COALESCE(daily_attendance.first_entry_at,EXCLUDED.first_entry_at),
                    status=CASE WHEN daily_attendance.first_entry_at IS NULL THEN EXCLUDED.status ELSE daily_attendance.status END,
                    calculated_at=now()""",(day,s['id'],occurred_at,status))
            else:
                c.execute("""INSERT INTO daily_attendance(attendance_date,student_id,last_exit_at,status)
                  VALUES(%s,%s,%s,'UNKNOWN') ON CONFLICT (attendance_date,student_id) DO UPDATE SET
                    last_exit_at=COALESCE(daily_attendance.last_exit_at,EXCLUDED.last_exit_at), calculated_at=now()""",(day,s['id'],occurred_at))
            if canonical:
                queued=c.execute("""INSERT INTO attendance_notifications(attendance_date,student_id,event_type)
                  VALUES(%s,%s,%s) ON CONFLICT (attendance_date,student_id,event_type) DO NOTHING RETURNING id""",(day,s['id'],event_type)).fetchone()
                notification_queued=queued is not None
        c.commit()
    return {'id':raw['id'] if raw else None,'accepted':True,'raw_recorded':raw is not None,'canonical_marking':canonical,'notification_queued':notification_queued}
